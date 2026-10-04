"""
Builds report/PA1_report.docx and exports it to PDF with Microsoft Word.

    cd ../pa1_starter/PA1_studentID_surname
    python report_figures.py                    # figures/report_fig*.png
    cd ../../report
    ../rl2026/Scripts/python.exe build_report.py

Tables 2 and 4 are filled from results_extra/extra_results.json.

Inline markup in the text below: **bold**, ~italic~ (a single * is literal, as in V*).
"""
from __future__ import annotations

import json
import pathlib
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

AUTHOR = "Seher Kanwal"
STUDENT_ID = "Student ID"          # <-- put your student ID here

HERE = pathlib.Path(__file__).resolve().parent
SUB = HERE.parent / "pa1_starter" / "PA1_studentID_surname"
FIGS = SUB / "figures"
DOCX = HERE / "PA1_report.docx"
PDF = SUB / "report.pdf"
EXTRA = json.loads((SUB / "results_extra" / "extra_results.json").read_text(encoding="utf-8"))


def num(x, d=1):
    """Number with a typographic minus."""
    return f"{x:.{d}f}".replace("-", "−")

NAVY = RGBColor(0x1F, 0x38, 0x64)
MUTED = RGBColor(0x52, 0x51, 0x4E)
BODY_PT, CAP_PT, TAB_PT = 10.5, 8.5, 8.5
TEXT_W = 6.9                        # inches between margins

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.75)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(BODY_PT)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.12


# ------------------------------------------------------------------ helpers
def runs(p, text, size=None, color=None):
    """Add text with **bold** / ~italic~ markup to paragraph p."""
    for tok in re.split(r"(\*\*.+?\*\*|~[^~]+?~)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("~"):
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def para(text, container=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, after=None, color=None):
    p = (container or doc).add_paragraph()
    p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return runs(p, text, size, color)


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = NAVY
    return p


def figure(name, width, caption, container=None):
    c = container or doc
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGS / f"report_{name}"), width=Inches(width))
    cap = para(caption, c, align=WD_ALIGN_PARAGRAPH.LEFT, size=CAP_PT, after=6, color=MUTED)
    return cap


def borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec)); el.set(qn("w:color"), "8A8A8A")
        else:
            el.set(qn("w:val"), "nil")
        tb.append(el)
    tcPr.append(tb)


def cell_margins(table, top=0, bottom=0, left=0, right=0):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); m.append(el)
    tblPr.append(m)


def layout(widths):
    """Borderless side-by-side layout table; returns its cells."""
    t = doc.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, left=60, right=60)
    for c, w in zip(t.rows[0].cells, widths):
        c.width = Inches(w)
        borders(c)
    return t.rows[0].cells


def tidy(cell):
    p = cell.paragraphs[0]                  # the empty default paragraph
    if not p.text and len(cell.paragraphs) > 1:
        p._element.getparent().remove(p._element)


def data_table(rows, widths, caption, container=None, bold_cells=()):
    c = container or doc
    t = c.add_table(rows=len(rows), cols=len(rows[0])) if c is not doc else         doc.add_table(rows=len(rows), cols=len(rows[0]))
    if c is not doc:                          # a cell appends an empty paragraph after a table
        extra = c.paragraphs[-1]
        extra._element.getparent().remove(extra._element)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t, top=15, bottom=15, left=70, right=70)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True      # keep the table on one page
            trPr = t.rows[i]._tr.get_or_add_trPr()
            if not trPr.findall(qn("w:cantSplit")):
                trPr.append(OxmlElement("w:cantSplit"))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            runs(p, val, size=TAB_PT)
            if i == 0 or (i, j) in bold_cells:
                for r in p.runs:
                    r.bold = True
            last = i == len(rows) - 1
            borders(cell, top=8 if i == 0 else None, bottom=8 if (i == 0 or last) else None)
    para(caption, c, align=WD_ALIGN_PARAGRAPH.LEFT, size=CAP_PT, after=6, color=MUTED)
    return t


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return runs(p, text)


# ------------------------------------------------------------------ title
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
r = p.add_run("PA1: From Bellman Updates to TD Control"); r.bold = True
r.font.size = Pt(17); r.font.color.rgb = NAVY
p = para(f"{AUTHOR}  ·  {STUDENT_ID}  ·  Reinforcement Learning (Graduate)  ·  October 2026",
         align=WD_ALIGN_PARAGRAPH.LEFT, size=10, after=6, color=MUTED)
pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
for k, v in (("val", "single"), ("sz", "6"), ("space", "4"), ("color", "1F3864")):
    b.set(qn(f"w:{k}"), v)
bdr.append(b); pPr.append(bdr)

para("All graded numbers are those in results.json (spec settings: α = 0.5, ε = 0.1, 3,000 episodes, "
     "400-step cap, seeds 0–4). Items marked **[supp.]** are supplementary: they come from "
     "extra_experiments.py, which reuses agent.epsilon_greedy and the same update rules (a check confirms "
     "it reproduces agent.sarsa / agent.q_learning bit for bit) and uses 20 seeds where stated. "
     "The report figures are drawn by report_figures.py; reproduction commands, seeds and versions are "
     "in README.md.", size=9, color=MUTED)

# ------------------------------------------------------------------ 1
heading("1  Policy iteration vs value iteration")
left, right = layout([3.35, 3.55])
figure("fig1_gridworld.png", 3.2,
       "**Figure 1.** RoomsGridWorld: V* (colour and value) and the greedy policy (arrows). "
       "PI and VI produce exactly this map.", left)
para("**Both methods reach the same optimal policy and values.** max |V_PI − V_VI| = 0 and every state "
     "gets the same action. The greedy path from S reaches G in 11 steps with undiscounted return 0 and "
     "V*(S) = −2.038. At S, ~up~ and ~right~ are exactly tied (Q = −2.038 for both; each route is 11 "
     "steps), and the lowest-index rule selects ~up~.", right)
para("**The counts differ because they count different things.** A PI round contains a complete policy "
     "evaluation; a VI sweep is a single backup of every state. Counted in sweeps, PI needs 3,276 against "
     "12 for VI, about 270 times more backups (Table 1).", right)
data_table([["", "Policy iteration", "Value iteration"],
            ["Iterations", "9 rounds", "12 sweeps"],
            ["Eval. sweeps per round", "540 (rounds 1–6), 12 (7–9)", "—"],
            ["Total sweeps", "3,276 + 9 improvements", "12"],
            ["V*(S), action at S", "−2.038, up", "−2.038, up"]],
           [1.25, 1.45, 0.85], "**Table 1.** DP summary; values identical to results.json.", right)
tidy(left); tidy(right)

para("**Why rounds 1–6 cost 540 sweeps each.** The initial all-~up~ policy drives most states into a wall "
     "or the top border forever, so their value is −1/(1−γ) = −20. Under a fixed policy the change per sweep "
     "shrinks only by γ = 0.95, and 0.95ᵏ⁻¹ < 10⁻¹² gives exactly k = 540. Each improvement re-routes only "
     "the 4–7 states next to already-good ones, so some state keeps looping until round 6. From round 7 "
     "every path terminates, and evaluating a deterministic terminating policy is exact after 12 sweeps. "
     "Each evaluation starts from V = 0; warm-starting from the previous V would lower the sweep total but "
     "not the 9 rounds, since every evaluation reaches the same fixed point.")
para("**Why VI needs only 12.** The max lets correct values spread outward from G one cell per sweep: the "
     "largest change is exactly 10·0.95ᵏ⁻¹ (10, 9.5, 9.03, …, 5.99). S is the farthest state, 11 steps "
     "away, so after 11 sweeps every value is exact and sweep 12 confirms Δ = 0. PI takes a few large, "
     "expensive steps (policy improvement is a Newton-like jump); VI takes many cheap ones. On this small "
     "deterministic grid VI is far cheaper.")

# ------------------------------------------------------------------ 2
heading("2  The paths learned by SARSA and Q-learning")
para("**Q-learning learns the shortest path along the cliff edge; SARSA learns a detour through the top "
     "rows** (Figure 2). Q-learning's greedy path runs along row 3 on every seed: 11 steps, return −11.0, "
     "**highest_row = 3.0**, identical to the DP optimum. SARSA's path climbs to row 0 (four seeds) or "
     "row 1 (one seed) before crossing: 15–17 steps, return −16.6 ± 0.8, **highest_row = 0.2**. With 20 "
     "seeds [supp.] the picture is the same: Q-learning 3.0 on every seed, SARSA rows 0–1 only.")
figure("fig2_paths.png", 6.0,
       "**Figure 2.** Greedy path of each seed (0–4), drawn slightly offset. Q-learning's path is identical "
       "on all seeds.")
para("**Why it leaves the edge.** Q-learning's target r + γ·max_a′ Q(s′,a′) assumes greedy behaviour from the next step on, so it "
     "estimates q*, the value of the optimal policy, whatever exploration is actually performed (Bellman "
     "optimality). SARSA's target r + γ·Q(s′,a′) uses the action the ε-greedy policy actually takes, so it "
     "estimates q_π of the ε-greedy policy itself (Bellman expectation). On row 3 every step has a "
     "0.1 × ¼ = 2.5 % chance of a random ~down~ move: −75 and a restart at S. SARSA's values price that risk "
     "in; Q-learning's do not. Figure 3 shows this: averaged over row 3 (columns 1–8), SARSA's value is "
     "−14.6 while Q-learning's is −5.35, exactly V*. (Q-learning matches V* on rows 2–3 and at S; only the "
     "rarely visited rows 0–1 are still optimistic, by up to 2.4, because Q starts at 0.)")
figure("fig3_value_maps.png", 6.6,
       "**Figure 3 [supp.].** Learned values max_a Q̄(s,a) and greedy action argmax_a Q̄(s,a) of the "
       "seed-averaged Q̄ (seeds 0–4). SARSA's values drop sharply next to the cliff; Q-learning's equal V* "
       "along the edge.")
para("**Why it goes as far as rows 0–1.** The risk explains leaving the edge, but not the full detour. "
     "SARSA's convergence target, the fixed point q*_ε of its expected update under ε-greedy behaviour "
     "(solved exactly from the model), keeps only one row of distance at ε = 0.1: its greedy path runs "
     "along row 2 with return −13 (Table 4). The extra climb to rows 0–1 (−16.6) comes from the constant "
     "α = 0.5 and the finite budget. On the same 20 seeds, α = 0.1 learns q*_ε's row-2 path on every seed "
     "(−13.0). With α = 0.5 each update moves Q halfway towards a single sampled target, so values next to "
     "the cliff stay noisy and are repeatedly knocked down by single fall-driven samples, and the greedy "
     "policy keeps more distance than the expected risk justifies. **highest_row therefore measures two "
     "things: the on-policy objective (row 3 → 2) and the step-size noise (row 2 → 0–1).**")

G = EXTRA["G"]


def learned(cell):
    rows = sorted(int(r) for r in cell["highest_row_counts"] if int(r) < 4)
    span = f"row {rows[0]}" if rows[0] == rows[-1] else f"rows {rows[0]}–{rows[-1]}"
    fail = cell["final_fail"]
    return f"{span} · {num(cell['greedy_mean_reached'])}" + (f" · **{fail} fail**" if fail else "")


t4 = [["ε", "q*_ε (exact)", "α = 0.5 (spec)", "α = 0.1", "α = 0.5 / (1 + n/50)"]]
for e in ["0.01", "0.1", "0.2", "0.3"]:
    fp = G["fixed_point"][e]
    t4.append([e, f"row {fp['highest_row']} · {num(fp['greedy_return'])}"] +
              [learned(G["sarsa"][a][e]) for a in ("alpha = 0.5 (spec)", "alpha = 0.1", "alpha = 0.5/(1+n/50)")])
data_table(t4, [0.45, 1.15, 1.75, 1.6, 1.75],
           "**Table 4 [supp.].** SARSA's convergence target vs what it learns (20 seeds, 3,000 episodes). "
           "Each cell: highest row of the greedy path · greedy return (mean over runs that reach G) · number of "
           "the 20 final greedy policies that never reach G. q*_ε is the exact fixed point under ε-greedy "
           "behaviour; n is the visit count of (s, a).")

# ------------------------------------------------------------------ 3
heading("3  The effect of ε")
para("We trained both algorithms with ε ∈ {0.01, 0.05, 0.1, 0.2, 0.3}, everything else fixed, with 5 seeds "
     "(results_sweep/) and again with 20 seeds [supp.]. The conclusions agree; Figure 4 and Table 2 report "
     "the 20-seed runs, which is why their ε = 0.1 row differs slightly from results.json. Besides training "
     "return and final greedy return we measured how many training episodes contain a cliff fall, and "
     "**greedy stability**: at the end of each of the last 500 episodes, does the current greedy policy "
     "reach G?")
figure("fig4_eps_sweep.png", 6.6,
       "**Figure 4 [supp.].** ε sweep, 20 seeds, mean ± 1 std across seeds; (a, b) over the last 500 "
       "training episodes. (c) Share of late-training greedy snapshots that reach G; the labels count "
       "runs whose final greedy policy never reaches G.")
C = EXTRA["C"]
t2 = [["ε", "Training return  S / Q", "Episodes with a fall  S / Q", "SARSA greedy reaches G",
       "SARSA greedy return: reached / all", "SARSA late snapshots reaching G"]]
for e in ["0.01", "0.05", "0.1", "0.2", "0.3"]:
    s, q = C[e]["sarsa"], C[e]["qlearning"]
    ts = f"{num(s['training_mean'])} ± {s['training_std']:.1f}"
    tq = f"{num(q['training_mean'])} ± {q['training_std']:.1f}"
    ts, tq = (f"**{ts}**", tq) if s["training_mean"] > q["training_mean"] else (ts, f"**{tq}**")
    t2.append([e, f"{ts} / {tq}", f"{100 * s['falls_mean']:.1f} % / {100 * q['falls_mean']:.1f} %",
               f"{s['greedy_reached']}/20", f"{num(s['greedy_mean_reached'])} / {num(s['greedy_mean_all'])}",
               f"{100 * s['snap_ok_mean']:.1f} %"])
data_table(t2, [0.4, 1.8, 1.35, 0.95, 1.3, 1.1],
           "**Table 2 [supp.].** 20 seeds, mean ± std; S = SARSA, Q = Q-learning; bold = better training "
           "return. Falls: share of the last 500 training episodes with at least one cliff transition "
           "(reward −75). SARSA's greedy return is given over the runs that reach G and over all 20 runs, "
           "where a policy that never reaches G scores −400 at the 400-step cap. Q-learning reaches G on "
           "20/20 seeds with return −11 and in 100 % of snapshots at every ε.")
para("**Exploration and training return.** Larger ε means more random moves and more falls, but the cost "
     "depends on where the agent walks. On the edge, Q-learning falls in 2.3 % → 46.5 % of episodes as ε "
     "goes from 0.01 to 0.3; SARSA, away from the edge, in 0.2 % → 15.6 %. Training return drops for both, "
     "much faster for Q-learning (−13.0 → −85.5 vs −16.5 → −46.6), and the ranking reverses: Q-learning "
     "trains better at ε = 0.01, they are level at 0.05, and SARSA is better from 0.1 on.")
para("**Final greedy evaluation.** Q-learning's greedy policy is optimal (−11) at every ε and seed, because "
     "its target does not depend on ε. SARSA's greedy path stays on the detour at every ε and becomes less "
     "reliable as ε grows: the share of late snapshots whose greedy policy reaches G falls from 98.6 % to "
     "60.6 %. At ε = 0.3, 5 of 20 final policies never reach G, which pulls the all-seed mean (capped "
     "at 400 steps) to −112.7 although the successful runs average −16.9.")
para("**These failures are not a bug.** The spec notes that a greedy policy which fails to reach G signals "
     "a bug, because this cannot happen ~at convergence~. Constant-α SARSA does not converge: its target "
     "samples a′ from the exploring policy (a random a′ near the cliff is worth about −100), and with "
     "α = 0.5 each noisy target moves Q halfway, so near-tied actions keep swapping. A greedy policy can "
     "then contain a loop. For example, at S (ε = 0.2, seed 3) Q(up) = −24.02 < Q(down) = −23.85, and "
     "~down~ bumps the border, so the greedy agent never leaves S; the ε-greedy behaviour escapes such "
     "loops, so training is unaffected. Table 4 confirms the diagnosis: q*_ε reaches G at every ε, and on "
     "the same 20 seeds α = 0.1 or a decaying α gives no failures at any ε (late-snapshot stability "
     "98–100 %). Q-learning, whose target has no sampled action, never fails, and all five graded SARSA "
     "runs (ε = 0.1, results.json) reach G.")
para("**Too little exploration also has a cost.** At ε = 0.01 the edge is optimal even for SARSA's own "
     "objective (q*_ε runs along row 3, Table 4), yet SARSA rarely finds it: "
     "1 of 20 runs with α = 0.1, none with α = 0.5. In a 20,000-episode run at ε = 0.01 [supp.] its greedy path only drifts from row 0 to rows 1–2. Early episodes, with Q initialised at 0, "
     "fall off the cliff often, so edge actions start with very negative values; with ε = 0.01 they are "
     "almost never retried, and those estimates are not corrected. In short, **ε trades training cost "
     "against discovery and, for SARSA, against greedy reliability.**")

# ------------------------------------------------------------------ 4
heading("4  Worse training return, better greedy return")
para("**Both facts hold because the two numbers measure different policies.** At the spec settings, "
     "Q-learning's training return is worse (−31.7 vs −23.2) while its greedy return is better (−11.0 vs "
     "−16.6). Training return measures the ε-greedy ~behaviour~ policy while it learns; greedy return "
     "measures the learned policy with exploration switched off (Figure 5).")
left, right = layout([3.4, 3.5])
figure("fig5_training_curves.png", 3.3,
       "**Figure 5.** Training return at the spec settings (ε = 0.1), mean ± 1 std over seeds 0–4, "
       "50-episode moving average, y-axis clipped at −100. Legend: raw last-500 means, as in results.json.",
       left)
para("To separate the two, we froze each learned Q and executed its policy without learning at "
     "different ε [supp.]. At ε = 0.1, SARSA's policy earns −20.9 and Q-learning's −31.7, the same as the "
     "DP-optimal policy executed ε-greedily (−31.8). At ε = 0 the order flips. The curves cross near "
     "ε ≈ 0.04 (Table 3, Figure 6).", right)
data_table([["Seeds 0–4", "SARSA", "Q-learning"],
            ["Training return (ε = 0.1, learning)", "**−23.2 ± 1.8**", "−31.7 ± 0.6"],
            ["Frozen policy at ε = 0.1 [supp.]", "**−20.9 ± 0.9**", "−31.7 ± 1.3"],
            ["Cliff falls / episode at ε = 0.1 [supp.]", "**0.03**", "0.24"],
            ["Greedy return (ε = 0)", "−16.6 ± 0.8", "**−11.0 ± 0.0**"],
            ["highest_row", "0.2", "3.0"]],
           [1.95, 0.75, 0.75],
           "**Table 3.** Mean ± std over seeds; frozen-policy rows average 2,000 episodes per seed.", right)
tidy(left); tidy(right)
figure("fig6_cross_eval.png", 6.6,
       "**Figure 6 [supp.].** Policies learned at ε = 0.1, frozen and executed at different ε. Q-learning's "
       "policy behaves exactly like the DP-optimal one (dashed).")
para("**What each algorithm is estimating.** Q-learning is off-policy: it estimates q*, the value of acting "
     "optimally with no future random moves. Its greedy policy is therefore optimal once exploration stops, "
     "but during training that policy is executed ε-greedily on the most dangerous route (0.24 falls per "
     "episode). SARSA is on-policy: it estimates q_π of the ε-greedy policy it actually runs, random moves "
     "included, so it tracks the policy that is best ~given that it keeps exploring~: a safer route with "
     "fewer falls and a better training return, but not the greedy optimum. Neither is wrong; the right "
     "choice depends on whether exploration continues when the policy is used. Decaying ε linearly from "
     "0.1 to 0 [supp.] confirms this: Q-learning's training return then reaches −11, equal to its greedy "
     "return, while SARSA's training and greedy returns both settle at −16.3 on the detour. With constant α "
     "and a finite budget, SARSA's early pessimism about the edge is locked in; reaching the optimum "
     "requires GLIE exploration ~and~ decreasing step sizes (Singh et al., 2000).")

# ------------------------------------------------------------------ 5
heading("5  Conclusions")
bullet("**DP:** PI and VI give identical V* and policies. VI needs 12 sweeps; PI needs 9 rounds but 3,276 "
       "evaluation sweeps, because early policies loop and (cold-started) evaluation converges only at rate γ.")
bullet("**On- vs off-policy:** Q-learning learns q* and the edge path (highest_row 3.0, return −11). SARSA "
       "learns the value of its ε-greedy behaviour and prices in the 2.5 % per-step fall risk; that alone "
       "moves it one row up (q*_ε: row 2, −13), and the constant α = 0.5 adds the rest (highest_row 0.2, "
       "return −16.6).")
bullet("**ε:** higher ε raises falls and lowers training return, far more for Q-learning, and makes SARSA's "
       "greedy policy less reliable (98.6 % → 60.6 %), a step-size effect that vanishes with α = 0.1. Very small ε reverses the training-return ranking and "
       "starves SARSA of the exploration it needs to revise early estimates.")
bullet("**Training vs greedy:** SARSA wins while exploring (−20.9 vs −31.7 at ε = 0.1), Q-learning wins "
       "without exploration (−11.0 vs −16.6); the crossover is near ε ≈ 0.04.")
para("**Limitations.** Graded results use 5 seeds (supplementary runs use 20). α is constant, so SARSA does "
     "not converge and its final greedy policy is a snapshot of a moving estimate. All findings come from "
     "one small deterministic environment. Further supplementary figures (PI/VI convergence trace, learning "
     "curves per ε, decaying-ε and 20,000-episode runs) are in figures/extra_*.png.", size=9, after=3)
para("**References.** R. S. Sutton & A. G. Barto (2018), ~Reinforcement Learning: An Introduction~, 2nd ed., "
     "Example 6.6. S. Singh, T. Jaakkola, M. L. Littman & C. Szepesvári (2000), Convergence results for "
     "single-step on-policy RL algorithms, ~Machine Learning~ 38:287–308.", size=9)

# ------------------------------------------------------------------ page numbers
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
    r._r.append(el)
r.font.size = Pt(8); r.font.color.rgb = MUTED

doc.save(DOCX)
print("wrote", DOCX)

# ------------------------------------------------------------------ PDF via Word
import win32com.client  # noqa: E402

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
try:
    d = word.Documents.Open(str(DOCX), ReadOnly=True)
    pages = d.ComputeStatistics(2)          # wdStatisticPages
    d.SaveAs2(str(PDF), FileFormat=17)      # wdFormatPDF
    d.Close(False)
finally:
    word.Quit()
print(f"wrote {PDF}  ({pages} pages)")
